#!/usr/bin/env python3
"""Read official CEFR-J workbooks without changing or exporting a workbook.

Requires openpyxl. Emits counts, sheet identities and version differences only;
does not assign Chinese grades, word senses, permissions or student difficulty.
"""
import argparse
import collections
import csv
import hashlib
import json
from pathlib import Path

import openpyxl


def identity(path):
    data = path.read_bytes()
    return {'bytes': len(data), 'sha256': hashlib.sha256(data).hexdigest()}


def word_rows(sheet):
    expected = ('headword', 'pos', 'CEFR', 'CoreInventory 1',
                'CoreInventory 2', 'Threshold')
    header = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    if header != expected:
        raise ValueError(f'{sheet.title}: unexpected columns {header}')
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    if any(not row[0] or not row[1] or row[2] not in ('A1', 'A2', 'B1', 'B2')
           for row in rows):
        raise ValueError(f'{sheet.title}: blank/invalid lexical row')
    return rows


def describe(rows):
    keys = [(r[0], r[1]) for r in rows]
    headwords = {r[0] for r in rows}
    levels = collections.defaultdict(set)
    for row in rows:
        levels[row[0]].add(row[2])
    return {
        'pos_entries': len(rows),
        'case_sensitive_headwords': len(headwords),
        'casefolded_headwords': len({x.casefold() for x in headwords}),
        'duplicate_headword_pos_entries': len(keys) - len(set(keys)),
        'entries_by_level': dict(sorted(collections.Counter(r[2] for r in rows).items())),
        'headwords_with_multiple_levels': sum(len(v) > 1 for v in levels.values()),
        'formula_cells': sum(isinstance(v, str) and v.startswith('=')
                             for row in rows for v in row),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--wordlist', type=Path, required=True)
    parser.add_argument('--grammar', type=Path, required=True)
    parser.add_argument('--previous-csv', type=Path)
    parser.add_argument('--previous-grammar-csv', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    result = {'method': 'read_only_openpyxl_raw_values_no_formula_recalculation',
              'tool_version': openpyxl.__version__, 'wordlist': identity(args.wordlist),
              'grammar': identity(args.grammar)}
    workbook = openpyxl.load_workbook(args.wordlist, read_only=True, data_only=False)
    sheets = {s.title: word_rows(s) for s in workbook if s.title != 'README'}
    result['wordlist']['sheets'] = {n: describe(r) for n, r in sheets.items()}
    for suffix in ('', '_sep'):
        all_rows = collections.Counter(sheets['ALL' + suffix])
        separate = collections.Counter(row for level in ('A1', 'A2', 'B1', 'B2')
                                        for row in sheets[level + suffix])
        if all_rows != separate:
            raise ValueError(f'ALL{suffix} differs from its four level sheets')
    expanded = collections.Counter((variant, *r[1:]) for r in sheets['ALL']
                                   for variant in r[0].split('/'))
    result['wordlist']['split_sheet_matches_slash_expansion'] = (
        expanded == collections.Counter(sheets['ALL_sep']))
    result['wordlist']['slash_expansion_only'] = [
        {'headword': r[0], 'pos': r[1], 'level': r[2], 'count': count}
        for r, count in (expanded - collections.Counter(sheets['ALL_sep'])).items()]
    result['wordlist']['split_sheet_only'] = [
        {'headword': r[0], 'pos': r[1], 'level': r[2], 'count': count}
        for r, count in (collections.Counter(sheets['ALL_sep']) - expanded).items()]
    result['wordlist']['split_sheet_matches_deduplicated_slash_expansion'] = (
        set(expanded) == set(sheets['ALL_sep']))
    if args.previous_csv:
        with args.previous_csv.open(newline='', encoding='utf-8-sig') as stream:
            reader = csv.DictReader(stream)
            previous = list(reader)
            headers = reader.fieldnames
        old = {(r['headword'], r['pos']): tuple(r[h] or None for h in headers)
               for r in previous}
        if len(old) != len(previous):
            raise ValueError('Previous CSV contains duplicate headword/POS keys')
        new = {(r[0], r[1]): r for r in sheets['ALL']}
        result['wordlist']['previous_csv_comparison'] = {
            **identity(args.previous_csv), 'rows': len(previous),
            'added_keys': [list(k) for k in sorted(new.keys() - old.keys())],
            'removed_keys': [list(k) for k in sorted(old.keys() - new.keys())],
            'changed_shared_keys': [
                {'headword': k[0], 'pos': k[1],
                 'changed_columns': [headers[i] for i, (a, b) in enumerate(zip(old[k], new[k]))
                                     if a != b]}
                for k in sorted(old.keys() & new.keys()) if old[k] != new[k]],
        }
    workbook.close()
    grammar = openpyxl.load_workbook(args.grammar, read_only=True, data_only=False)
    result['grammar']['sheets'] = [{'name': s.title, 'rows': s.max_row,
                                     'columns': s.max_column,
                                     'version_cell_A1': s.cell(1, 1).value}
                                    for s in grammar]
    for name in ('ITEM LIST', '教員版'):
        items = [r for r in grammar[name].iter_rows(min_row=3, values_only=True)
                 if r[0] is not None and r[1] is not None]
        ids = [str(r[0]) for r in items]
        base_ids = [int(x.split('-')[0]) for x in ids]
        result['grammar'][name] = {
            'item_rows': len(items), 'distinct_ids': len(set(ids)),
            'compound_ids': sum('-' in x for x in ids),
            'base_id_min': min(base_ids), 'base_id_max': max(base_ids),
            'missing_base_ids_inside_range': sorted(
                set(range(min(base_ids), max(base_ids) + 1)) - set(base_ids)),
        }
        if name == '教員版':
            result['grammar'][name]['levels'] = dict(sorted(
                collections.Counter(str(r[4]) for r in items).items()))
    if args.previous_grammar_csv:
        with args.previous_grammar_csv.open(newline='', encoding='utf-8-sig') as stream:
            old_grammar = list(csv.DictReader(stream))
        previous_ids = {r['ID'] for r in old_grammar}
        current_ids = {str(r[0]) for r in grammar['ITEM LIST'].iter_rows(
            min_row=3, values_only=True) if r[0] is not None and r[1] is not None}
        result['grammar']['previous_csv_id_comparison'] = {
            **identity(args.previous_grammar_csv), 'rows': len(old_grammar),
            'added_ids': sorted(current_ids - previous_ids),
            'removed_ids': sorted(previous_ids - current_ids),
            'limitation': 'ID comparison only; no equivalence of regexes, CEFR labels or definitions inferred',
        }
    grammar.close()
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'word_entries': result['wordlist']['sheets']['ALL']['pos_entries'],
                      'word_headwords': result['wordlist']['sheets']['ALL']['case_sensitive_headwords'],
                      'grammar_sheets': len(result['grammar']['sheets'])}))


if __name__ == '__main__':
    main()
